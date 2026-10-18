"""Turn class slides (PDF / PPTX) into a runnable lab script."""
import io
import json
import re

from core import llm

GUI_HINTS = ("gui", "클릭", "누른다", "선택", "메뉴")
COMMAND = re.compile(r"^(cd|ffmpeg|colmap|python3?|pip3?|sudo|apt(-get)?|conda|git|mkdir|cp|mv|ls|bash|sh|wget|curl|unzip|jupyter)\b")


def extract_text(filename: str, data: bytes) -> list[str]:
    """One string per slide/page."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return [p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages]
    if name.endswith(".pptx"):
        from pptx import Presentation
        return ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in Presentation(io.BytesIO(data)).slides]
    return [data.decode("utf-8", "replace")]


def _command_lines(pages):
    """Command-looking lines with their slide numbers, joined when a command wraps onto the next line."""
    found = []
    for i, page in enumerate(pages, 1):
        lines = [l.strip().lstrip("-•●").strip() for l in page.splitlines()]
        for j, line in enumerate(lines):
            if COMMAND.match(line):
                nxt = lines[j + 1] if j + 1 < len(lines) else ""
                if line.count('"') % 2 and nxt:  # an unclosed quote: the command continues on the next line
                    line = f"{line} {nxt}"
                found.append((i, line))
    return found


def rule_lab(pages):
    """Fallback without the LLM: keep concrete commands, skip templates ({...}) and GUI launches."""
    script, manual, workdir = [], [], "~/lab"
    commands = _command_lines(pages)
    first_cmd_page = commands[0][0] if commands else 0
    for page, cmd in commands:
        cmd = re.sub(r"\s*\(.*\)\s*$", "", cmd).strip()
        if cmd.startswith("cd "):
            workdir = cmd[3:].strip()
            continue
        if "{" in cmd:
            continue  # a template line; the slide gives a concrete example right after
        if cmd.endswith(" gui") or " gui " in cmd:
            manual.append(f"슬라이드 {page}: `{cmd}` (GUI 실행, 자동 리허설 불가)")
            continue
        script.append(cmd)
    # Slides about GUI tools with no commands on them: list them as manual steps, grouped by tool.
    for tool in ("COLMAP", "CloudCompare", "Omniverse"):
        # A slide "about" a tool names it in its title (near the top), after the hands-on part starts.
        hits = [i for i, p in enumerate(pages, 1) if i > first_cmd_page and tool.lower() in p.strip()[:80].lower()
                and not any(COMMAND.match(l.strip().lstrip("-•●").strip()) for l in p.splitlines())]
        if len(hits) > 1:
            manual.append(f"슬라이드 {hits[0]}~{hits[-1]}: {tool} 화면 조작 {len(hits)}장 (GUI, 자동 리허설 불가)")
    return {"workdir": workdir, "script": script, "manual_steps": manual, "by": "rules"}


PROMPT = """These are the slides of a hands-on class. Students follow them on their own PCs.
Extract the steps a student types in a terminal into a bash script that Dry Run can run on test machines.

Slides:
{slides}

Rules:
- "workdir": the folder students work in (from a `cd` step), e.g. "~/Desktop/1018recon".
- "script": the terminal commands, in order, WITHOUT the `cd` (Dry Run runs inside workdir). Use the slide's concrete example
  values instead of {{placeholder}} templates. Keep commands exactly as a student would type them; do not add fixes.
- Leave out commands that only open a GUI (e.g. `colmap gui`) and put them, with every click-only step, in "manual_steps"
  as short Korean sentences that start with the slide number ("슬라이드 12: ...").
Reply with ONLY JSON: {{"workdir": "...", "script": ["cmd", ...], "manual_steps": ["...", ...]}}"""


def to_lab(pages):
    lab = rule_lab(pages)
    if lab["script"]:
        return lab  # commands found on the slides: deterministic and exact
    slides = "\n".join(f"[slide {i}] {p.strip()[:600]}" for i, p in enumerate(pages, 1) if p.strip())
    try:
        raw = llm.chat([{"role": "user", "content": PROMPT.format(slides=slides[:12000])}], temperature=0.1, max_tokens=1500)
        data = json.loads(raw[raw.find("{"): raw.rfind("}") + 1])
        if not data.get("script"):
            raise ValueError("no commands")
        data["by"] = "llm"
    except Exception as e:
        data = rule_lab(pages)
        data["llm_error"] = str(e)[:200]
    return data


def lesson_script(lab) -> str:
    return "\n".join(["set -e", f"cd {lab['workdir']}", *lab["script"]]) + "\n"
